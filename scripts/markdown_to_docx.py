from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "二十一协作区微信小程序开发需求文档.md"
DEFAULT_OUTPUT = ROOT / "二十一协作区微信小程序开发需求文档.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10.5)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)

    for name, size, color in [
        ("Title", 26, (18, 59, 50)),
        ("Heading 1", 18, (18, 59, 50)),
        ("Heading 2", 15, (18, 59, 50)),
        ("Heading 3", 12, (154, 112, 44)),
    ]:
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)


def add_cover(document, title, subtitle_text, version):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(90)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(18, 59, 50)

    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_paragraph.paragraph_format.space_before = Pt(16)
    run = subtitle_paragraph.add_run(subtitle_text)
    run.bold = True
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(177, 132, 61)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(50)
    run = note.add_run("四队联合公益服务与发展档案 · 内部资料")
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(130, 130, 130)

    info = document.add_table(rows=4, cols=2)
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.style = "Table Grid"
    values = [
        ("项目负责人", ""),
        ("协作区负责人", ""),
        ("计划上线时间", ""),
        ("文档版本", version),
    ]
    for row, (label, value) in zip(info.rows, values):
        set_cell_text(row.cells[0], label, bold=True, color=(18, 59, 50))
        set_cell_shading(row.cells[0], "E8EEE9")
        set_cell_text(row.cells[1], value)

    document.add_page_break()


def add_fill_area(document, lines=3):
    for _ in range(lines):
        paragraph = document.add_paragraph("________________________________________________________________________________")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.runs[0].font.color.rgb = RGBColor(170, 170, 170)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(document, rows):
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        for col_index in range(width):
            value = values[col_index] if col_index < len(values) else ""
            set_cell_text(
                table.rows[row_index].cells[col_index],
                value,
                bold=row_index == 0,
                color=(255, 255, 255) if row_index == 0 else None,
            )
            if row_index == 0:
                set_cell_shading(table.rows[row_index].cells[col_index], "123B32")
            elif row_index % 2 == 0:
                set_cell_shading(table.rows[row_index].cells[col_index], "F5F2E9")
    document.add_paragraph()


def add_body(document, markdown):
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines = []

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if line.startswith("```"):
            if in_code:
                meaningful = [item for item in code_lines if item.strip()]
                if meaningful:
                    for item in meaningful:
                        paragraph = document.add_paragraph()
                        run = paragraph.add_run(item)
                        run.font.name = "微软雅黑"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                        run.font.size = Pt(10)
                add_fill_area(document, 3)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not line:
            index += 1
            continue

        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1 and title == "二十一协作区微信小程序开发需求文档":
                index += 1
                continue
            document.add_heading(title, level=level)
            index += 1
            continue

        if line.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(line.lstrip("> ").strip())
            run.italic = True
            run.font.color.rgb = RGBColor(105, 115, 110)
            index += 1
            continue

        checkbox = re.match(r"^-\s+\[([ xX])\]\s+(.+)$", line)
        if checkbox:
            symbol = "☒" if checkbox.group(1).lower() == "x" else "☐"
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{symbol} {checkbox.group(2)}")
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            document.add_paragraph(bullet.group(1), style="List Bullet")
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            document.add_paragraph(numbered.group(1), style="List Number")
            index += 1
            continue

        paragraph = document.add_paragraph()
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", line)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.color.rgb = RGBColor(154, 112, 44)
            else:
                paragraph.add_run(part)
        index += 1


def add_footer(document):
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("二十一协作区微信小程序开发需求文档")
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(145, 145, 145)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--title", default="二十一协作区")
    parser.add_argument("--subtitle", default="微信小程序开发需求文档")
    parser.add_argument("--version", default="V1.0")
    args = parser.parse_args()

    source = args.source.resolve()
    output = args.output.resolve()
    markdown = source.read_text(encoding="utf-8")
    document = Document()
    configure_document(document)
    add_cover(document, args.title, args.subtitle, args.version)
    add_body(document, markdown)
    add_footer(document)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
